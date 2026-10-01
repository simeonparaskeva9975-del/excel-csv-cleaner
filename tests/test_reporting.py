"""Tests for report-table construction and Excel export."""

from __future__ import annotations

import pandas as pd
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from cleaner_app.analysis import analyze_dataset
from cleaner_app.cleaning import FOUNDATION_ACTIONS, clean_dataset
from cleaner_app.file_io import LoadedDataset, load_dataset
from cleaner_app.models import CleaningAction

from cleaner_app.reporting import (
    build_cleaning_log_table,
    build_summary_table,
    export_cleaned_csv,
    export_excel_report,
)

from cleaner_app.reporting import (
    build_cleaning_log_table,
    build_summary_table,
    export_excel_report,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = (
    PROJECT_ROOT
    / "sample_data"
    / "messy_sales_data.xlsx"
)


class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.loaded_dataset = load_dataset(XLSX_PATH)
        cls.before_profile = analyze_dataset(cls.loaded_dataset)

        all_actions = FOUNDATION_ACTIONS | {
            CleaningAction.STANDARDIZE_TEXT_CASE,
            CleaningAction.CONVERT_NUMERIC,
            CleaningAction.CONVERT_DATES,
        }

        cls.cleaning_result = clean_dataset(
            cls.loaded_dataset,
            actions=all_actions,
            text_case_columns=(
                "Product Category",
                "City ",
                "Sales Channel",
                "Order Status",
            ),
            numeric_columns=(
                "Quantity",
                "Unit Price EUR",
                "Discount Percent",
            ),
            date_columns=("Order Date",),
        )

        cleaned_dataset = LoadedDataset(
            data=cls.cleaning_result.data,
            source_path=Path("cleaned_sales_data.xlsx"),
            file_type="Excel",
            file_size_bytes=0,
        )

        cls.after_profile = analyze_dataset(cleaned_dataset)

    def test_cleaned_csv_can_be_exported(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = (
                Path(temporary_directory)
                / "cleaned_sales_data.csv"
            )

            returned_path = export_cleaned_csv(
                output_path,
                self.cleaning_result,
            )

            self.assertEqual(returned_path, output_path)
            self.assertTrue(output_path.exists())

            exported_data = pd.read_csv(output_path)

            self.assertEqual(exported_data.shape, (120, 12))
            self.assertEqual(
                exported_data.columns.tolist(),
                self.cleaning_result.data.columns.tolist(),
            )
            self.assertNotIn(
                "Unnamed: 0",
                exported_data.columns,
            )

    def test_excel_report_contains_meaningful_charts(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = (
                Path(temporary_directory)
                / "report_with_charts.xlsx"
            )

            export_excel_report(
                output_path,
                self.cleaning_result,
                self.before_profile,
                self.after_profile,
            )

            workbook = load_workbook(output_path)

            charts_sheet = workbook["Charts"]

            self.assertEqual(
                charts_sheet["A1"].value,
                "Report Charts",
            )
            self.assertGreaterEqual(
                len(charts_sheet._charts),
                2,
            )
            self.assertEqual(
                charts_sheet["A3"].value,
                "Column",
            )
            self.assertEqual(
                charts_sheet["B3"].value,
                "Missing Values",
            )
            self.assertFalse(
                charts_sheet.sheet_view.showGridLines
            )

            workbook.close()

    def test_csv_export_rejects_wrong_extension(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            invalid_path = (
                Path(temporary_directory)
                / "cleaned_sales_data.txt"
            )

            with self.assertRaises(ValueError):
                export_cleaned_csv(
                    invalid_path,
                    self.cleaning_result,
                )

    def test_excel_report_has_professional_formatting(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = (
                Path(temporary_directory)
                / "formatted_report.xlsx"
            )

            export_excel_report(
                output_path,
                self.cleaning_result,
                self.before_profile,
                self.after_profile,
            )

            workbook = load_workbook(output_path)

            cleaned_sheet = workbook["Cleaned Data"]
            summary_sheet = workbook["Data Quality Summary"]
            log_sheet = workbook["Cleaning Log"]

            self.assertIsNone(cleaned_sheet.auto_filter.ref)
            self.assertIsNone(log_sheet.auto_filter.ref)
            self.assertEqual(cleaned_sheet.freeze_panes, "A2")
            self.assertIn(
                "CleanedDataTable",
                cleaned_sheet.tables,
            )
            self.assertGreater(
                cleaned_sheet.column_dimensions["A"].width,
                5,
            )

            self.assertIn(
                "A1:D1",
                {
                    str(merged_range)
                    for merged_range in summary_sheet.merged_cells.ranges
                },
            )
            self.assertTrue(summary_sheet["A1"].font.bold)
            self.assertEqual(summary_sheet.freeze_panes, "A8")
            self.assertFalse(
                summary_sheet.sheet_view.showGridLines
            )

            longest_email_length = max(
                len(str(value))
                for value in self.cleaning_result.data[
                    "email_address"
                ]
                if value is not None
            )

            self.assertGreaterEqual(
                cleaned_sheet.column_dimensions["C"].width,
                longest_email_length + 2,
            )

            self.assertEqual(log_sheet.freeze_panes, "A2")
            self.assertIn(
                "CleaningLogTable",
                log_sheet.tables,
            )

            workbook.close()

    def test_summary_table_contains_before_after_metrics(
        self,
    ) -> None:
        summary = build_summary_table(
            self.before_profile,
            self.after_profile,
        )

        self.assertEqual(
            summary.columns.tolist(),
            ["Metric", "Before", "After", "Difference"],
        )
        self.assertEqual(len(summary), 14)

        metrics = summary.set_index("Metric")

        self.assertEqual(
            metrics.loc["Exact Duplicate Rows", "Before"],
            8,
        )
        self.assertEqual(
            metrics.loc["Exact Duplicate Rows", "After"],
            0,
        )
        self.assertEqual(
            metrics.loc["Fully Blank Columns", "Before"],
            1,
        )
        self.assertEqual(
            metrics.loc["Fully Blank Columns", "After"],
            0,
        )
        self.assertEqual(
            metrics.loc["Outer Whitespace Cells", "Before"],
            25,
        )
        self.assertEqual(
            metrics.loc["Outer Whitespace Cells", "After"],
            0,
        )
        self.assertEqual(
            metrics.loc["Case-Variant Groups", "Before"],
            4,
        )
        self.assertEqual(
            metrics.loc["Case-Variant Groups", "After"],
            0,
        )
        self.assertEqual(
            metrics.loc["Ambiguous Date Values", "Before"],
            15,
        )
        self.assertEqual(
            metrics.loc["Ambiguous Date Values", "After"],
            14,
        )
        self.assertEqual(
            metrics.loc["Invalid Date Values", "After"],
            2,
        )

    def test_cleaning_log_table_is_complete(self) -> None:
        log_table = build_cleaning_log_table(
            self.cleaning_result
        )

        self.assertEqual(len(log_table), 33)

        self.assertEqual(
            log_table.columns.tolist(),
            [
                "Timestamp (UTC)",
                "Action",
                "Column",
                "Description",
                "Affected Rows",
                "Affected Cells",
                "Before",
                "After",
            ],
        )

        self.assertEqual(
            int(log_table["Affected Rows"].sum()),
            11,
        )

        actions = set(log_table["Action"])

        self.assertIn("remove_duplicates", actions)
        self.assertIn("trim_whitespace", actions)
        self.assertIn("standardize_text_case", actions)
        self.assertIn("convert_numeric", actions)
        self.assertIn("convert_dates", actions)

    def test_excel_report_contains_expected_sheets_and_data(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = (
                Path(temporary_directory)
                / "cleaned_sales_report.xlsx"
            )

            returned_path = export_excel_report(
                output_path,
                self.cleaning_result,
                self.before_profile,
                self.after_profile,
            )

            self.assertEqual(returned_path, output_path)
            self.assertTrue(output_path.exists())

            workbook = load_workbook(
                output_path,
                data_only=False,
            )

            self.assertEqual(
                workbook.sheetnames,
                [
                    "Cleaned Data",
                    "Data Quality Summary",
                    "Cleaning Log",
                    "Charts"
                ],
            )

            cleaned_sheet = workbook["Cleaned Data"]

            self.assertEqual(cleaned_sheet.max_row, 121)
            self.assertEqual(cleaned_sheet.max_column, 12)
            self.assertEqual(
                cleaned_sheet["A1"].value,
                "order_id",
            )

            summary_sheet = workbook[
                "Data Quality Summary"
            ]

            self.assertEqual(
                summary_sheet["A1"].value,
                "Data Quality Summary",
            )
            self.assertEqual(
                summary_sheet["B3"].value,
                "messy_sales_data.xlsx",
            )
            self.assertEqual(
                summary_sheet["B4"].value,
                131,
            )
            self.assertEqual(
                summary_sheet["B5"].value,
                120,
            )

            log_sheet = workbook["Cleaning Log"]

            self.assertEqual(log_sheet.max_row, 34)
            self.assertEqual(
                log_sheet["A1"].value,
                "Timestamp (UTC)",
            )
            self.assertEqual(
                log_sheet["B1"].value,
                "Action",
            )

            workbook.close()


if __name__ == "__main__":
    unittest.main()