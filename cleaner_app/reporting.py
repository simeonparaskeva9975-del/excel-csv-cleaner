"""Report-table construction and Excel export utilities."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from datetime import date, datetime

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference

DARK_BLUE = "17324D"
ACCENT_BLUE = "1E88A8"
LIGHT_BLUE = "DCEAF2"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"

import pandas as pd
from openpyxl import Workbook

from cleaner_app.models import CleaningResult, DatasetProfile


def profile_metric_values(
    profile: DatasetProfile,
) -> dict[str, int]:
    """Convert a profile into report-friendly quality metrics."""

    case_variant_groups = sum(
        len(groups)
        for groups in profile.case_variants_by_column.values()
    )

    numeric_text_values = sum(
        finding.numeric_text_count
        for finding in profile.numeric_findings.values()
    )

    invalid_numeric_values = sum(
        finding.invalid_count
        for finding in profile.numeric_findings.values()
    )

    ambiguous_date_values = sum(
        finding.ambiguous_text_count
        for finding in profile.date_findings.values()
    )

    invalid_date_values = sum(
        finding.invalid_count
        for finding in profile.date_findings.values()
    )

    negative_numeric_values = sum(
        finding.negative_count
        for finding in (
            profile.suspicious_numeric_findings.values()
        )
    )

    numeric_outliers = sum(
        finding.outlier_count
        for finding in (
            profile.suspicious_numeric_findings.values()
        )
    )

    return {
        "Total Rows": profile.row_count,
        "Total Columns": profile.column_count,
        "Exact Duplicate Rows": profile.duplicate_rows,
        "Fully Blank Rows": profile.blank_rows,
        "Fully Blank Columns": len(profile.blank_columns),
        "Missing Values": sum(
            profile.missing_by_column.values()
        ),
        "Outer Whitespace Cells": sum(
            profile.whitespace_cells_by_column.values()
        ),
        "Case-Variant Groups": case_variant_groups,
        "Numeric Text Values": numeric_text_values,
        "Invalid Numeric Values": invalid_numeric_values,
        "Ambiguous Date Values": ambiguous_date_values,
        "Invalid Date Values": invalid_date_values,
        "Negative Numeric Values": negative_numeric_values,
        "Potential Numeric Outliers": numeric_outliers,
    }


def build_summary_table(
    before_profile: DatasetProfile,
    after_profile: DatasetProfile,
) -> pd.DataFrame:
    """Build the before-and-after data-quality summary."""

    before_metrics = profile_metric_values(before_profile)
    after_metrics = profile_metric_values(after_profile)

    rows = []

    for metric, before_value in before_metrics.items():
        after_value = after_metrics[metric]

        rows.append(
            {
                "Metric": metric,
                "Before": before_value,
                "After": after_value,
                "Difference": after_value - before_value,
            }
        )

    return pd.DataFrame(rows)


def build_cleaning_log_table(
    cleaning_result: CleaningResult,
) -> pd.DataFrame:
    """Convert cleaning log entries into a report table."""

    rows = []

    for entry in cleaning_result.log_entries:
        rows.append(
            {
                "Timestamp (UTC)": (
                    entry.created_at_utc.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                ),
                "Action": entry.action.value,
                "Column": entry.column or "",
                "Description": entry.description,
                "Affected Rows": entry.affected_rows,
                "Affected Cells": entry.affected_cells,
                "Before": (
                    ""
                    if entry.before_value is None
                    else str(entry.before_value)
                ),
                "After": (
                    ""
                    if entry.after_value is None
                    else str(entry.after_value)
                ),
            }
        )

    return pd.DataFrame(
        rows,
        columns=[
            "Timestamp (UTC)",
            "Action",
            "Column",
            "Description",
            "Affected Rows",
            "Affected Cells",
            "Before",
            "After",
        ],
    )


def excel_safe_value(value: Any) -> Any:
    """Convert pandas and NumPy values into Excel-safe values."""

    if value is None:
        return None

    try:
        if bool(pd.isna(value)):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if hasattr(value, "item"):
        try:
            return value.item()
        except (TypeError, ValueError):
            pass

    return value


def write_dataframe(
    worksheet: Any,
    data: pd.DataFrame,
    *,
    start_row: int = 1,
    start_column: int = 1,
) -> None:
    """Write a DataFrame, including headers, to a worksheet."""

    for column_offset, column_name in enumerate(
        data.columns,
        start=start_column,
    ):
        worksheet.cell(
            row=start_row,
            column=column_offset,
            value=str(column_name),
        )

    for row_offset, row_values in enumerate(
        data.itertuples(index=False, name=None),
        start=start_row + 1,
    ):
        for column_offset, value in enumerate(
            row_values,
            start=start_column,
        ):
            worksheet.cell(
                row=row_offset,
                column=column_offset,
                value=excel_safe_value(value),
            )

def style_header_row(
    worksheet: Any,
    row_number: int,
    first_column: int,
    last_column: int,
) -> None:
    """Apply the standard report header style."""

    for column_number in range(
        first_column,
        last_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=DARK_BLUE,
        )
        cell.font = Font(
            color=WHITE,
            bold=True,
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,

        )

    worksheet.row_dimensions[row_number].height = 22


def set_automatic_column_widths(
    worksheet: Any,
    *,
    maximum_width: int = 35,
) -> None:
    """Set readable widths while preventing extremely wide columns."""

    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):
        longest_value = 0

        for row_number in range(
            1,
            worksheet.max_row + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if value is None:
                continue

            longest_value = max(
                longest_value,
                len(str(value)),
            )

        column_letter = get_column_letter(column_number)

        worksheet.column_dimensions[column_letter].width = min(
            max(longest_value + 2, 10),
            maximum_width,
        )


def add_excel_table(
    worksheet: Any,
    *,
    table_name: str,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
) -> None:
    """Add a professional striped Excel table."""

    if end_row <= start_row:
        return

    reference = (
        f"{get_column_letter(start_column)}{start_row}:"
        f"{get_column_letter(end_column)}{end_row}"
    )

    table = Table(
        displayName=table_name,
        ref=reference,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(table)


def style_cleaned_data_sheet(worksheet: Any) -> None:
    """Style the cleaned-data worksheet."""

    worksheet.freeze_panes = "A2"
    worksheet.sheet_properties.tabColor = ACCENT_BLUE
    worksheet.sheet_view.zoomScale = 90

    style_header_row(
        worksheet,
        row_number=1,
        first_column=1,
        last_column=worksheet.max_column,
    )

    add_excel_table(
        worksheet,
        table_name="CleanedDataTable",
        start_row=1,
        start_column=1,
        end_row=worksheet.max_row,
        end_column=worksheet.max_column,
    )

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
            )

            if isinstance(cell.value, (datetime, date)):
                cell.number_format = "yyyy-mm-dd"

    set_automatic_column_widths(
        worksheet,
        maximum_width=60,
    )


def style_summary_sheet(
    worksheet: Any,
    summary_table: pd.DataFrame,
) -> None:
    """Style the data-quality summary worksheet."""

    worksheet.merge_cells("A1:D1")
    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor=DARK_BLUE,
    )
    worksheet["A1"].font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 30

    for row_number in range(3, 6):
        label_cell = worksheet.cell(
            row=row_number,
            column=1,
        )
        label_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=LIGHT_BLUE,
        )
        label_cell.font = Font(bold=True)

    style_header_row(
        worksheet,
        row_number=7,
        first_column=1,
        last_column=4,
    )

    first_metric_row = 8
    last_metric_row = (
        first_metric_row + len(summary_table) - 1
    )

    # Apply status colors only to quality metrics, not dimensions.
    for row_number in range(
        first_metric_row + 2,
        last_metric_row + 1,
    ):
        after_cell = worksheet.cell(
            row=row_number,
            column=3,
        )
        difference_cell = worksheet.cell(
            row=row_number,
            column=4,
        )

        if after_cell.value == 0:
            after_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=LIGHT_GREEN,
            )

        if (
            isinstance(difference_cell.value, (int, float))
            and difference_cell.value < 0
        ):
            difference_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=LIGHT_GREEN,
            )
        elif (
            isinstance(difference_cell.value, (int, float))
            and difference_cell.value > 0
        ):
            difference_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=LIGHT_RED,
            )

    worksheet.freeze_panes = "A8"
    worksheet.sheet_properties.tabColor = DARK_BLUE
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 100

    set_automatic_column_widths(
        worksheet,
        maximum_width=70,
    )


def style_cleaning_log_sheet(worksheet: Any) -> None:
    """Style the cleaning-log worksheet."""

    worksheet.freeze_panes = "A2"
    worksheet.sheet_properties.tabColor = "70AD47"
    worksheet.sheet_view.zoomScale = 85

    style_header_row(
        worksheet,
        row_number=1,
        first_column=1,
        last_column=worksheet.max_column,
    )

    add_excel_table(
        worksheet,
        table_name="CleaningLogTable",
        start_row=1,
        start_column=1,
        end_row=worksheet.max_row,
        end_column=worksheet.max_column,
    )

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column == 4,
            )

    set_automatic_column_widths(
        worksheet,
        maximum_width=70,
    )

def choose_category_counts(
    data: pd.DataFrame,
) -> tuple[str, pd.Series] | None:
    """Choose one meaningful low-cardinality text column."""

    candidates: list[tuple[int, int, str, pd.Series]] = []

    for column in data.columns:
        text_values = [
            value.strip()
            for value in data[column]
            if isinstance(value, str) and value.strip()
        ]

        if len(text_values) < 10:
            continue

        counts = pd.Series(text_values).value_counts()
        unique_count = len(counts)

        if unique_count < 2 or unique_count > 10:
            continue

        if unique_count / len(text_values) > 0.25:
            continue

        candidates.append(
            (
                unique_count,
                -len(text_values),
                str(column),
                counts,
            )
        )

    if not candidates:
        return None

    candidates.sort(
        key=lambda item: (
            item[0],
            item[1],
            item[2],
        )
    )

    _, _, column, counts = candidates[0]

    return column, counts


def create_charts_sheet(
    workbook: Workbook,
    cleaning_result: CleaningResult,
    after_profile: DatasetProfile,
) -> None:
    """Create report charts only when suitable data exists."""

    worksheet = workbook.create_sheet("Charts")

    worksheet["A1"] = "Report Charts"
    worksheet.merge_cells("A1:J1")
    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor=DARK_BLUE,
    )
    worksheet["A1"].font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    worksheet["A1"].alignment = Alignment(
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 30

    worksheet.sheet_properties.tabColor = "FFC000"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90

    missing_values = [
        (column, count)
        for column, count in (
            after_profile.missing_by_column.items()
        )
        if count > 0
    ]

    next_section_row = 3

    if missing_values:
        worksheet.cell(
            row=next_section_row,
            column=1,
            value="Column",
        )
        worksheet.cell(
            row=next_section_row,
            column=2,
            value="Missing Values",
        )

        style_header_row(
            worksheet,
            row_number=next_section_row,
            first_column=1,
            last_column=2,
        )

        for row_offset, (column, count) in enumerate(
            missing_values,
            start=next_section_row + 1,
        ):
            worksheet.cell(
                row=row_offset,
                column=1,
                value=column,
            )
            worksheet.cell(
                row=row_offset,
                column=2,
                value=count,
            )

        missing_end_row = (
            next_section_row + len(missing_values)
        )

        missing_chart = BarChart()
        missing_chart.type = "bar"
        missing_chart.style = 10
        missing_chart.title = "Remaining Missing Values"
        missing_chart.x_axis.title = "Missing Values"
        missing_chart.y_axis.title = "Column"
        missing_chart.height = 7
        missing_chart.width = 12
        missing_chart.legend = None

        missing_data = Reference(
            worksheet,
            min_col=2,
            min_row=next_section_row,
            max_row=missing_end_row,
        )
        missing_categories = Reference(
            worksheet,
            min_col=1,
            min_row=next_section_row + 1,
            max_row=missing_end_row,
        )

        missing_chart.add_data(
            missing_data,
            titles_from_data=True,
        )
        missing_chart.set_categories(
            missing_categories
        )

        worksheet.add_chart(
            missing_chart,
            "D3",
        )

        next_section_row = max(
            missing_end_row + 3,
            20,
        )

    category_result = choose_category_counts(
        cleaning_result.data
    )

    if category_result is not None:
        category_column, category_counts = category_result

        worksheet.cell(
            row=next_section_row,
            column=1,
            value=category_column,
        )
        worksheet.cell(
            row=next_section_row,
            column=2,
            value="Count",
        )

        style_header_row(
            worksheet,
            row_number=next_section_row,
            first_column=1,
            last_column=2,
        )

        for row_offset, (category, count) in enumerate(
            category_counts.items(),
            start=next_section_row + 1,
        ):
            worksheet.cell(
                row=row_offset,
                column=1,
                value=category,
            )
            worksheet.cell(
                row=row_offset,
                column=2,
                value=int(count),
            )

        category_end_row = (
            next_section_row + len(category_counts)
        )

        category_chart = BarChart()
        category_chart.type = "col"
        category_chart.style = 10
        category_chart.title = (
            f"Category Counts: {category_column}"
        )
        category_chart.x_axis.title = category_column
        category_chart.y_axis.title = "Count"
        category_chart.height = 7
        category_chart.width = 12
        category_chart.legend = None

        category_data = Reference(
            worksheet,
            min_col=2,
            min_row=next_section_row,
            max_row=category_end_row,
        )
        category_categories = Reference(
            worksheet,
            min_col=1,
            min_row=next_section_row + 1,
            max_row=category_end_row,
        )

        category_chart.add_data(
            category_data,
            titles_from_data=True,
        )
        category_chart.set_categories(
            category_categories
        )

        worksheet.add_chart(
            category_chart,
            f"D{next_section_row}",
        )

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 16

def export_excel_report(
    output_path: str | Path,
    cleaning_result: CleaningResult,
    before_profile: DatasetProfile,
    after_profile: DatasetProfile,
) -> Path:
    """Export cleaned data, summary, and cleaning log to Excel."""

    destination = Path(output_path)

    if destination.suffix.lower() != ".xlsx":
        raise ValueError(
            "Excel report output must use the .xlsx extension."
        )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    summary_table = build_summary_table(
        before_profile,
        after_profile,
    )
    log_table = build_cleaning_log_table(
        cleaning_result
    )

    workbook = Workbook()

    cleaned_sheet = workbook.active
    cleaned_sheet.title = "Cleaned Data"

    summary_sheet = workbook.create_sheet(
        "Data Quality Summary"
    )
    log_sheet = workbook.create_sheet(
        "Cleaning Log"
    )

    write_dataframe(
        cleaned_sheet,
        cleaning_result.data,
    )

    summary_sheet["A1"] = "Data Quality Summary"
    summary_sheet["A3"] = "Source File"
    summary_sheet["B3"] = cleaning_result.source_name
    summary_sheet["A4"] = "Original Rows"
    summary_sheet["B4"] = (
        cleaning_result.original_row_count
    )
    summary_sheet["A5"] = "Cleaned Rows"
    summary_sheet["B5"] = (
        cleaning_result.cleaned_row_count
    )

    write_dataframe(
        summary_sheet,
        summary_table,
        start_row=7,
    )

    write_dataframe(
        log_sheet,
        log_table,
    )

    style_cleaned_data_sheet(cleaned_sheet)
    style_summary_sheet(summary_sheet, summary_table)
    style_cleaning_log_sheet(log_sheet)

    create_charts_sheet(
        workbook,
        cleaning_result,
        after_profile,
    )

    workbook.properties.creator = "Excel/CSV Cleaner"
    workbook.properties.title = "Data Cleaning Report"
    workbook.properties.subject = (
        "Cleaned dataset and data-quality summary"
    )

    workbook.save(destination)

    return destination

def export_cleaned_csv(
    output_path: str | Path,
    cleaning_result: CleaningResult,
) -> Path:
    """Export only the cleaned dataset as a UTF-8 CSV file."""

    destination = Path(output_path)

    if destination.suffix.lower() != ".csv":
        raise ValueError(
            "Cleaned CSV output must use the .csv extension."
        )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    cleaning_result.data.to_csv(
        destination,
        index=False,
        encoding="utf-8-sig",
        date_format="%Y-%m-%d",
    )

    return destination


    workbook.save(destination)

    return destination